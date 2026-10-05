from __future__ import annotations

import json
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from app.connectors.clockify import ClockifyClient
from app.reporting.engine import generate_report
from app.reporting.profile import DEFAULT_PROFILE


def resolve_profile(template_path: Path, report_type: str) -> dict:
    workbook = load_workbook(template_path, read_only=True)
    cloud9_monthly = False
    atm_helper_monthly = False
    try:
        visible_sheets = [sheet.title for sheet in workbook.worksheets if sheet.sheet_state == "visible"]
        active_sheet = workbook.active.title if workbook.active and workbook.active.sheet_state == "visible" else None
        if report_type == "monthly" and "Timesheet" in visible_sheets:
            cloud9_headers = (
                "Sr. No.", "Resource Name", "Date", "Time", "Date", "Time", "Duration (H:MM)", "Task",
                "Description / Comments",
            )
            sheet = workbook["Timesheet"]
            headers = tuple(str(sheet.cell(7, column).value or "").strip() for column in range(1, 10))
            cloud9_monthly = headers == cloud9_headers
        atm_helper_monthly = (
            report_type == "monthly"
            and "Auto Timesheet" in visible_sheets
            and "Clockify Raw" in workbook.sheetnames
            and "Clockify Helper" in workbook.sheetnames
        )
    finally:
        workbook.close()

    candidate_sheets = ([active_sheet] if active_sheet else []) + [name for name in visible_sheets if name != active_sheet]
    profile_path: Path | None = None
    profile_dir = Path(__file__).resolve().parent.parent / "config" / "profiles"
    if cloud9_monthly:
        profile_path = profile_dir / "cloud9_monthly.json"
    elif atm_helper_monthly:
        profile_path = profile_dir / "atm_link_helper_monthly.json"
    if profile_path:
        return json.loads(profile_path.read_text(encoding="utf-8"))
    if "Detailed Report" in candidate_sheets:
        profile_path = profile_dir / "detailed_weekly.json"
    elif "Auto Timesheet" in candidate_sheets:
        profile_path = profile_dir / "atm_link_monthly.json"
    if profile_path:
        return json.loads(profile_path.read_text(encoding="utf-8"))
    if DEFAULT_PROFILE["sheet"] in candidate_sheets:
        return DEFAULT_PROFILE
    raise RuntimeError(
        "No report mapping matches the active template sheet. Visible sheets: "
        + repr(candidate_sheets)
    )


def main() -> None:
    payload = json.load(sys.stdin)
    start = datetime.fromisoformat(payload["date_start"])
    end = datetime.fromisoformat(payload["date_end"])
    client = ClockifyClient(
        api_key=payload["api_key"],
        api_base=payload.get("api_base", "https://api.clockify.me/api/v1"),
        reports_base=payload.get("reports_base", "https://reports.api.clockify.me/v1"),
    )
    entries = client.detailed_entries(
        workspace_id=payload["workspace_id"],
        date_start=start,
        date_end=end,
        timezone=payload["timezone"],
        project_ids=payload.get("project_ids") or None,
    )
    if not entries:
        raise RuntimeError("Clockify returned no entries for the test report period")

    entry_owners = sorted({entry.user_name.strip() for entry in entries if entry.user_name.strip()})
    employee_name = ", ".join(entry_owners) if entry_owners else payload["employee_name"]

    template_path = Path(payload["template_path"])
    profile = resolve_profile(template_path, payload["report_type"])
    result = generate_report(
        template_path=template_path,
        output_path=Path(payload["output_path"]),
        entries=entries,
        profile=profile,
        employee_name=employee_name,
        project_name=payload["project_name"],
        period_start=start,
        period_end=end,
    )
    print(json.dumps({"validation": result.model_dump(), "employee_name": employee_name}))


if __name__ == "__main__":
    main()
