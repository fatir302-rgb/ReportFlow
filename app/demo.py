from __future__ import annotations

from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.models import NormalizedTimeEntry


def demo_entries() -> list[NormalizedTimeEntry]:
    starts = [
        datetime(2026, 8, 17, 20, 0, tzinfo=timezone.utc),
        datetime(2026, 8, 18, 20, 15, tzinfo=timezone.utc),
        datetime(2026, 8, 19, 19, 50, tzinfo=timezone.utc),
        datetime(2026, 8, 20, 20, 5, tzinfo=timezone.utc),
        datetime(2026, 8, 21, 20, 0, tzinfo=timezone.utc),
    ]
    durations = [8*3600, 7*3600+45*60, 8*3600+10*60, 8*3600, 7*3600+50*60]
    descriptions = [
        "Website maintenance and reporting",
        "Client updates and development",
        "Design implementation",
        "Testing and fixes",
        "Weekly wrap-up and deployment",
    ]
    rows: list[NormalizedTimeEntry] = []
    for i, start in enumerate(starts):
        end = start + timedelta(seconds=durations[i])
        rows.append(NormalizedTimeEntry(
            id=f"demo-{i+1}",
            user_name="Demo User",
            project_name="Demo Project",
            description=descriptions[i],
            start=start,
            end=end,
            duration_seconds=durations[i],
        ))
    return rows


def create_demo_template(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Timesheet"
    ws["A1"] = "WEEKLY PROJECT TIMESHEET"
    ws["A1"].font = Font(size=16, bold=True)
    ws.merge_cells("A1:F1")
    ws["A1"].alignment = Alignment(horizontal="center")

    ws["A2"] = "Employee"
    ws["A3"] = "Project"
    ws["A4"] = "Period"
    ws["E5"] = "Total"
    for cell in ["A2", "A3", "A4", "E5"]:
        ws[cell].font = Font(bold=True)

    headers = ["Date", "Project", "Description", "Start", "End", "Duration"]
    for col, value in enumerate(headers, start=1):
        cell = ws.cell(7, col, value)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.alignment = Alignment(horizontal="center")

    # Style row used by the engine as a reusable data-row pattern.
    for col in range(1, 7):
        ws.cell(8, col).alignment = Alignment(vertical="top", wrap_text=True)

    widths = {"A": 16, "B": 22, "C": 45, "D": 14, "E": 14, "F": 14}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
