from __future__ import annotations

from copy import copy
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any
import re

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.worksheet.cell_range import CellRange
from openpyxl.worksheet.worksheet import Worksheet

from app.models import NormalizedTimeEntry, ValidationResult


class ReportGenerationError(RuntimeError):
    pass


def _copy_row_style(ws: Worksheet, source_row: int, target_row: int, max_col: int) -> None:
    if source_row == target_row:
        return
    for col in range(1, max_col + 1):
        src = ws.cell(source_row, col)
        dst = ws.cell(target_row, col)
        if src.has_style:
            dst._style = copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        if src.font:
            dst.font = copy(src.font)
        if src.fill:
            dst.fill = copy(src.fill)
        if src.border:
            dst.border = copy(src.border)
        if src.alignment:
            dst.alignment = copy(src.alignment)
        if src.protection:
            dst.protection = copy(src.protection)
    if source_row in ws.row_dimensions:
        ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height


def _set(ws: Worksheet, cell: str | None, value: Any) -> None:
    if cell:
        ws[cell] = value


def _sort_entries(entries: list[NormalizedTimeEntry], profile: dict[str, Any]) -> list[NormalizedTimeEntry]:
    sort = profile.get("sort") or {}
    field = sort.get("field", "start")
    reverse = str(sort.get("direction", "asc")).lower() == "desc"
    if field not in {"start", "end", "project_name", "description"}:
        field = "start"
    return sorted(entries, key=lambda item: getattr(item, field), reverse=reverse)


def _entry_value(key: str, entry: NormalizedTimeEntry, project_name: str) -> Any:
    if key == "date" or key == "start_date":
        return entry.start.replace(tzinfo=None).date()
    if key == "end_date":
        return entry.end.replace(tzinfo=None).date()
    if key == "project":
        return entry.project_name or project_name
    if key == "description":
        return entry.description
    if key == "task":
        return entry.task_name
    if key == "start" or key == "start_time":
        return entry.start.replace(tzinfo=None).time()
    if key == "end" or key == "end_time":
        return entry.end.replace(tzinfo=None).time()
    if key == "duration":
        return entry.duration_seconds / 86400
    if key == "duration_decimal":
        return entry.duration_seconds / 3600
    if key == "user":
        return entry.user_name
    raise ReportGenerationError(f"Unsupported tabular field mapping: {key}")


def _apply_format(cell, key: str, formats: dict[str, str]) -> None:
    if key in {"date", "start_date", "end_date"}:
        cell.number_format = formats.get("date", "yyyy-mm-dd")
    elif key in {"start", "end", "start_time", "end_time"}:
        cell.number_format = formats.get("time", "h:mm AM/PM")
    elif key == "duration":
        cell.number_format = formats.get("duration", "[h]:mm")
    elif key == "duration_decimal":
        cell.number_format = formats.get("decimal", "0.00")


def _generate_tabular(
    workbook,
    entries: list[NormalizedTimeEntry],
    profile: dict[str, Any],
    employee_name: str,
    project_name: str,
    period_start: datetime,
    period_end: datetime,
) -> int:
    sheet_name = profile.get("sheet", "Timesheet")
    if sheet_name not in workbook.sheetnames:
        raise ReportGenerationError(f"Sheet '{sheet_name}' not found. Existing sheets: {workbook.sheetnames}")
    ws = workbook[sheet_name]

    data_start = int(profile.get("data_start_row", 8))
    style_source = int(profile.get("style_source_row", data_start))
    clear_extra = int(profile.get("clear_extra_rows", 20))
    columns: dict[str, str] = profile["columns"]
    metadata = profile.get("metadata", {})
    formats = profile.get("formats", {})
    metadata_sheet = workbook[profile.get("metadata_sheet", sheet_name)]
    period_sheet = workbook[profile.get("period_sheet", profile.get("metadata_sheet", sheet_name))]

    _set(metadata_sheet, metadata.get("employee_name"), employee_name)
    _set(metadata_sheet, metadata.get("project_name"), project_name)
    if metadata.get("period_start"):
        _set(metadata_sheet, metadata["period_start"], period_start)
        metadata_sheet[metadata["period_start"]].number_format = formats.get("date", "yyyy-mm-dd")
    if metadata.get("period_end"):
        display_end = period_end + timedelta(days=int(profile.get("period_end_offset_days", 0)))
        _set(metadata_sheet, metadata["period_end"], display_end)
        metadata_sheet[metadata["period_end"]].number_format = formats.get("date", "yyyy-mm-dd")
    period_format = profile.get("period_format", "{start:%d %b %Y} - {end:%d %b %Y}")
    _set(period_sheet, metadata.get("period"), period_format.format(start=period_start, end=period_end))

    title = profile.get("title") or {}
    if title.get("cell"):
        week_number = period_start.isocalendar().week
        title_text = str(title.get("template", "{project_name} - Week {week_number} - {employee_name}"))
        ws[title["cell"]] = title_text.format(
            employee_name=employee_name,
            project_name=project_name,
            week_number=week_number,
            period_start=period_start,
            period_end=period_end,
        )

    ordered_entries = _sort_entries(entries, profile)
    max_clear_row = max(ws.max_row, data_start + max(len(ordered_entries), 1) + clear_extra)
    for row in range(data_start, max_clear_row + 1):
        for column in columns.values():
            ws[f"{column}{row}"].value = None

    max_col = min(ws.max_column, int(profile.get("style_copy_columns", ws.max_column)))
    generated_seconds = 0
    for index, entry in enumerate(ordered_entries):
        row = data_start + index
        _copy_row_style(ws, style_source, row, max_col)
        for key, column in columns.items():
            cell = ws[f"{column}{row}"]
            cell.value = index + 1 if key == "sequence" else _entry_value(key, entry, project_name)
            _apply_format(cell, key, formats)
        if not _is_break(entry):
            generated_seconds += entry.duration_seconds

    total_cell = metadata.get("total_hours")
    if total_cell:
        ws[total_cell] = generated_seconds / 86400
        ws[total_cell].number_format = formats.get("duration", "[h]:mm")

    return generated_seconds


def _generate_import_linked_template(
    workbook,
    entries: list[NormalizedTimeEntry],
    profile: dict[str, Any],
    employee_name: str,
    project_name: str,
    period_start: datetime,
    period_end: datetime,
) -> tuple[int, str | None]:
    import_sheet = profile.get("import_sheet", "Clockify Import")
    if import_sheet not in workbook.sheetnames:
        raise ReportGenerationError(f"Import sheet '{import_sheet}' not found. Existing sheets: {workbook.sheetnames}")
    import_ws = workbook[import_sheet]
    metadata_sheet_name = profile.get("metadata_sheet", profile.get("sheet", "Timesheet"))
    if metadata_sheet_name not in workbook.sheetnames:
        raise ReportGenerationError(f"Metadata sheet '{metadata_sheet_name}' not found. Existing sheets: {workbook.sheetnames}")
    metadata_ws = workbook[metadata_sheet_name]
    period_sheet_name = profile.get("period_sheet", import_sheet)
    if period_sheet_name not in workbook.sheetnames:
        raise ReportGenerationError(f"Period sheet '{period_sheet_name}' not found. Existing sheets: {workbook.sheetnames}")
    period_ws = workbook[period_sheet_name]

    metadata = profile.get("metadata", {})
    _set(metadata_ws, metadata.get("employee_name"), employee_name)
    _set(metadata_ws, metadata.get("project_name"), project_name)
    display_end = period_end + timedelta(days=int(profile.get("period_end_offset_days", 0)))
    period_format = profile.get("period_format", "{start:%d %b %Y} - {end:%d %b %Y}")
    _set(period_ws, metadata.get("period"), period_format.format(start=period_start, end=display_end))

    start_row = int(profile.get("import_start_row", 7))
    max_entries = int(profile.get("import_max_rows", 1000))
    columns: dict[str, str] = profile["import_columns"]
    if len(entries) > max_entries:
        raise ReportGenerationError(f"Report has {len(entries)} entries but template supports only {max_entries}")

    for row in range(start_row, start_row + max_entries):
        for column in columns.values():
            import_ws[f"{column}{row}"].value = None

    generated_seconds = 0
    ordered_entries = _sort_entries(entries, profile)
    for row, entry in enumerate(ordered_entries, start_row):
        start = entry.start.replace(tzinfo=None)
        end = entry.end.replace(tzinfo=None)
        values = {
            "date": start.date(),
            "description": entry.description or entry.task_name,
            "project": entry.project_name or project_name,
            "duration": entry.duration_seconds / 86400,
            "start": start,
            "end": end,
            "ends_next_day": "Yes" if end.date() > start.date() else "No",
            "user": entry.user_name or employee_name,
        }
        for key, column in columns.items():
            import_ws[f"{column}{row}"] = values[key]
        generated_seconds += entry.duration_seconds

    total_column = profile.get("total_formula_column")
    if total_column and ordered_entries:
        total_cell = f"{total_column}{start_row + len(ordered_entries) + int(profile.get('linked_row_offset', 2))}"
        ws[total_cell] = generated_seconds / 86400
        ws[total_cell].number_format = profile.get("formats", {}).get("duration", "[h]:mm")
    return generated_seconds, None

def _generate_clockify_helper_calendar(
    workbook,
    entries: list[NormalizedTimeEntry],
    profile: dict[str, Any],
    employee_name: str,
    period_start: datetime,
) -> tuple[int, str | None]:
    sheet_name = profile.get("sheet", "Auto Timesheet")
    raw_sheet_name = profile.get("raw_sheet", "Clockify Raw")
    required = (sheet_name, raw_sheet_name, "Clockify Helper")
    missing = [name for name in required if name not in workbook.sheetnames]
    if missing:
        raise ReportGenerationError(f"ATM Link formula sheets missing: {', '.join(missing)}")

    ws = workbook[sheet_name]
    raw_ws = workbook[raw_sheet_name]
    raw_start = int(profile.get("raw_start_row", 2))
    raw_max_row = int(profile.get("raw_max_row", 1000))
    columns: dict[str, str] = profile["raw_columns"]
    max_entries = raw_max_row - raw_start + 1
    ordered = sorted(entries, key=lambda entry: entry.start)
    if len(ordered) > max_entries:
        raise ReportGenerationError(f"Report has {len(ordered)} entries but ATM template supports only {max_entries}")

    _set(ws, profile.get("resource_name_cell"), employee_name)
    month_anchor = profile.get("month_anchor_cell")
    if month_anchor:
        ws[month_anchor] = period_start.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
        ws[month_anchor].number_format = "mmmm yyyy"

    # Clear every source field so sample values in unused columns cannot leak into the report.
    for row in range(raw_start, raw_max_row + 1):
        for column in range(1, 16):
            raw_ws.cell(row, column).value = None

    generated_seconds = 0
    formats = profile.get("formats", {})
    style_source_row = int(profile.get("style_source_row", raw_start))
    entries_by_day: dict[Any, list[NormalizedTimeEntry]] = {}
    for row, entry in enumerate(ordered, raw_start):
        if row != style_source_row:
            _copy_row_style(raw_ws, style_source_row, row, raw_ws.max_column)
        start = entry.start.replace(tzinfo=None)
        end = entry.end.replace(tzinfo=None)
        duration_hours = entry.duration_seconds / 3600
        duration_whole_minutes = int(entry.duration_seconds // 60)
        values = {
            "project": entry.project_name,
            "description": entry.description,
            "task": entry.task_name,
            "user": entry.user_name or employee_name,
            "start_date": start.date(),
            "start_time": start.strftime("%I:%M %p").removeprefix("0"),
            "end_date": end.date(),
            "end_time": end.strftime("%I:%M %p").removeprefix("0"),
            "duration_text": f"{duration_whole_minutes // 60}:{duration_whole_minutes % 60:02d}",
            "duration_decimal": duration_hours,
        }
        for key, column in columns.items():
            cell = raw_ws[f"{column}{row}"]
            cell.value = values[key]
            if key in {"start_date", "end_date"}:
                cell.number_format = formats.get("date", "yyyy-mm-dd")
            elif key == "duration_decimal":
                cell.number_format = formats.get("decimal", "0.00")
        entries_by_day.setdefault(start.date(), []).append(entry)
        if not _is_break(entry):
            generated_seconds += entry.duration_seconds

    first_day_row = int(profile.get("first_day_row", 5))
    base_detail_rows = int(profile.get("base_detail_rows", 5))
    block_size = int(profile.get("day_block_size", base_detail_rows + 2))
    # Remove old formulas/values from the visible calendar without changing its styles.
    for row in range(first_day_row, ws.max_row):
        for column in "DEFGH":
            cell = ws[f"{column}{row}"]
            if not isinstance(cell, MergedCell):
                cell.value = None

    extra_rows_by_day: dict[int, int] = {}
    next_month = (period_start.replace(day=28) + timedelta(days=4)).replace(day=1)
    days_in_month = (next_month - timedelta(days=1)).day
    for day in range(1, 32):
        if day <= days_in_month:
            current_day = period_start.replace(day=day).date()
            count = len(entries_by_day.get(current_day, []))
        else:
            count = 0
        extra_rows_by_day[day] = max(0, count - base_detail_rows)

    # Insert additional detail lines from bottom upward, matching the template's cleanup macro.
    for day in range(31, 0, -1):
        extra = extra_rows_by_day[day]
        if not extra:
            continue
        summary_row = first_day_row + (day - 1) * block_size
        insert_row = summary_row + base_detail_rows + 1
        ws.insert_rows(insert_row, extra)
        for merged_range in list(ws.merged_cells.ranges):
            if merged_range.min_row >= insert_row:
                ws.merged_cells.remove(merged_range)
                ws.merged_cells.add(CellRange(
                    min_col=merged_range.min_col,
                    min_row=merged_range.min_row + extra,
                    max_col=merged_range.max_col,
                    max_row=merged_range.max_row + extra,
                ))
            elif merged_range.max_row >= insert_row:
                ws.merged_cells.remove(merged_range)
                ws.merged_cells.add(CellRange(
                    min_col=merged_range.min_col,
                    min_row=merged_range.min_row,
                    max_col=merged_range.max_col,
                    max_row=merged_range.max_row + extra,
                ))
        style_row = insert_row - 1
        for offset in range(extra):
            target_row = insert_row + offset
            _copy_row_style(ws, style_row, target_row, 8)
            ws.row_dimensions[target_row].height = ws.row_dimensions[style_row].height

    # Write final values into the template's styled calendar cells; don't rely on Excel formulas or its cleanup macro.
    extra_before = 0
    daily_total_cells: list[str] = []
    for day in range(1, 32):
        summary_row = first_day_row + (day - 1) * block_size + extra_before
        day_date = period_start.replace(day=day).date() if day <= days_in_month else None
        day_entries = sorted(entries_by_day.get(day_date, []), key=lambda item: item.start) if day_date else []
        slot_count = base_detail_rows + extra_rows_by_day[day]
        if summary_row != first_day_row:
            for column in range(4, 9):
                source_cell = ws.cell(first_day_row, column)
                target_cell = ws.cell(summary_row, column)
                if source_cell.has_style:
                    target_cell._style = copy(source_cell._style)
            ws.row_dimensions[summary_row].height = ws.row_dimensions[first_day_row].height
        date_cell = ws[f"D{summary_row}"]
        date_cell.value = f"{day_date.strftime('%B')} {day}, {day_date.year}" if day_date else None
        date_cell.number_format = "mmmm d, yyyy"
        work_entries = [entry for entry in day_entries if not _is_break(entry)]
        is_weekend = bool(day_date and day_date.weekday() >= 5)
        ws[f"G{summary_row}"] = "Work Total" if work_entries or not is_weekend else None
        ws[f"H{summary_row}"] = (
            sum(entry.duration_seconds for entry in work_entries) / 86400
            if work_entries or not is_weekend
            else "Weekend"
        )
        ws[f"H{summary_row}"].number_format = formats.get("duration", "[h]:mm")
        daily_total_cells.append(f"H{summary_row}")

        for index in range(1, slot_count + 1):
            detail_row = summary_row + index
            if index > len(day_entries):
                ws.row_dimensions[detail_row].hidden = True
                for column in "DEFGH":
                    ws[f"{column}{detail_row}"].value = None
                continue
            ws.row_dimensions[detail_row].hidden = False
            entry = day_entries[index - 1]
            start = entry.start.replace(tzinfo=None)
            end = entry.end.replace(tzinfo=None)
            ws[f"D{detail_row}"] = start.time()
            ws[f"E{detail_row}"] = end.time()
            ws[f"F{detail_row}"] = "Break" if _is_break(entry) else (entry.description or entry.task_name or entry.project_name)
            ws[f"G{detail_row}"] = "Break" if _is_break(entry) else "Work"
            ws[f"H{detail_row}"] = entry.duration_seconds / 86400
            for column in "DE":
                ws[f"{column}{detail_row}"].number_format = formats.get("time", "h:mm AM/PM")
            ws[f"H{detail_row}"].number_format = formats.get("duration", "[h]:mm")
        extra_before += extra_rows_by_day[day]

    total_cell = profile.get("grand_total_cell")
    generated_total_cell: str | None = None
    if total_cell:
        match = re.fullmatch(r"([A-Z]+)(\d+)", total_cell.upper())
        if not match:
            raise ReportGenerationError(f"Invalid ATM grand-total cell: {total_cell}")
        total_row = int(match.group(2)) + sum(extra_rows_by_day.values())
        generated_total_cell = f"{match.group(1)}{total_row}"
        ws[generated_total_cell] = f"=SUM({','.join(daily_total_cells)})"
        ws[generated_total_cell].number_format = formats.get("duration", "[h]:mm")

    ws["E3"] = period_start.replace(day=1).strftime("%B %Y")

    return generated_seconds, generated_total_cell


def _task_text(entry: NormalizedTimeEntry) -> str:
    return entry.description or entry.task_name or entry.project_name or ""


def _is_break(entry: NormalizedTimeEntry) -> bool:
    return (entry.project_name or "").strip().lower() == "break"


def _mentions_sheet(text: Any, sheet_names: list[str]) -> bool:
    """True if a formula/reference string points at any of the given sheet names."""
    if not isinstance(text, str) or not text:
        return False
    lowered = text.lower()
    return any(f"{name.lower()}!" in lowered or f"'{name.lower()}'!" in lowered for name in sheet_names)


def _scrub_removed_sheet_references(workbook, removed_sheets: list[str]) -> None:
    """Drop workbook content that still points at sheets ReportFlow deleted.

    Excel reports "We found a problem with some content" when defined names,
    conditional formatting, data validation or hyperlinks reference a sheet that
    no longer exists in the file.
    """
    if not removed_sheets:
        return

    for name, defined in list(workbook.defined_names.items()):
        if _mentions_sheet(getattr(defined, "attr_text", ""), removed_sheets):
            del workbook.defined_names[name]

    for ws in workbook.worksheets:
        for name, defined in list(ws.defined_names.items()):
            if _mentions_sheet(getattr(defined, "attr_text", ""), removed_sheets):
                del ws.defined_names[name]

        kept_rules = []
        for cf in ws.conditional_formatting:
            for rule in cf.rules:
                if not any(_mentions_sheet(formula, removed_sheets) for formula in (rule.formula or [])):
                    kept_rules.append((cf.sqref, rule))
        if len(kept_rules) != sum(len(cf.rules) for cf in ws.conditional_formatting):
            ws.conditional_formatting = ConditionalFormattingList()
            for sqref, rule in kept_rules:
                ws.conditional_formatting.add(str(sqref), rule)

        if ws.data_validations and ws.data_validations.dataValidation:
            ws.data_validations.dataValidation = [
                dv
                for dv in ws.data_validations.dataValidation
                if not (
                    _mentions_sheet(dv.formula1, removed_sheets)
                    or _mentions_sheet(dv.formula2, removed_sheets)
                )
            ]

        for cell in ws._cells.values():
            link = cell.hyperlink
            if link is not None and _mentions_sheet(getattr(link, "location", None), removed_sheets):
                cell.hyperlink = None


def _finalize_presentation_only(
    workbook,
    profile: dict[str, Any],
    generated_formula_cells: set[tuple[str, str]] | None = None,
) -> None:
    generated_formula_cells = generated_formula_cells or set()
    output_sheets = profile.get("output_sheets") or [profile.get("sheet", workbook.active.title)]
    missing = [name for name in output_sheets if name not in workbook.sheetnames]
    if missing:
        raise ReportGenerationError(f"Output sheet(s) not found: {', '.join(missing)}")
    removed_sheets: list[str] = []
    for sheet_name in list(workbook.sheetnames):
        if sheet_name not in output_sheets:
            workbook.remove(workbook[sheet_name])
            removed_sheets.append(sheet_name)
    _scrub_removed_sheet_references(workbook, removed_sheets)
    for sheet_name in output_sheets:
        ws = workbook[sheet_name]
        ws.sheet_state = "visible"
        for cell in ws._cells.values():
            if (
                (cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("=")))
                and (sheet_name, cell.coordinate) not in generated_formula_cells
            ):
                cell.value = None
    workbook.active = workbook.sheetnames.index(output_sheets[0])
    # The template's tab bar may start at a later sheet (firstSheet=1). After other
    # sheets are removed that index no longer exists, and Excel reports the file as corrupt.
    for view in workbook.views:
        view.firstSheet = 0
        view.activeTab = workbook.sheetnames.index(output_sheets[0])


def _write_raw_sheet(workbook, entries: list[NormalizedTimeEntry], profile: dict[str, Any]) -> None:
    raw_sheet_name = profile.get("raw_sheet")
    if not raw_sheet_name or raw_sheet_name not in workbook.sheetnames:
        return
    ws = workbook[raw_sheet_name]
    start_row = int(profile.get("raw_start_row", 2))
    max_rows = int(profile.get("raw_max_rows", 1000))

    headers = [
        "Project", "Client", "Description", "Task", "User", "Group", "Email", "Tags",
        "Billable", "Start Date", "Start Time", "End Date", "End Time", "Duration (h)", "Duration (decimal)"
    ]
    for col, header in enumerate(headers, 1):
        if not ws.cell(1, col).value:
            ws.cell(1, col).value = header

    for row in range(start_row, max_rows + 1):
        for col in range(1, 16):
            ws.cell(row, col).value = None

    ordered = sorted(entries, key=lambda item: item.start)
    for index, entry in enumerate(ordered, start=start_row):
        if index > max_rows:
            raise ReportGenerationError(f"Raw sheet capacity exceeded ({max_rows - start_row + 1} entries)")
        values = [
            entry.project_name,
            None,
            entry.description,
            entry.task_name,
            entry.user_name,
            None,
            None,
            None,
            None,
            entry.start.replace(tzinfo=None).date(),
            entry.start.replace(tzinfo=None).time(),
            entry.end.replace(tzinfo=None).date(),
            entry.end.replace(tzinfo=None).time(),
            entry.duration_seconds / 86400,
            entry.duration_seconds / 3600,
        ]
        for col, value in enumerate(values, 1):
            ws.cell(index, col).value = value
        ws.cell(index, 10).number_format = "yyyy-mm-dd"
        ws.cell(index, 11).number_format = "h:mm AM/PM"
        ws.cell(index, 12).number_format = "yyyy-mm-dd"
        ws.cell(index, 13).number_format = "h:mm AM/PM"
        ws.cell(index, 14).number_format = "[h]:mm"
        ws.cell(index, 15).number_format = "0.00"


def _generate_calendar_timesheet(
    workbook,
    entries: list[NormalizedTimeEntry],
    profile: dict[str, Any],
    employee_name: str,
    period_start: datetime,
) -> int:
    sheet_name = profile.get("sheet", "Auto Timesheet")
    if sheet_name not in workbook.sheetnames:
        raise ReportGenerationError(f"Sheet '{sheet_name}' not found. Existing sheets: {workbook.sheetnames}")
    ws = workbook[sheet_name]

    first_day_row = int(profile.get("first_day_row", 5))
    block_size = int(profile.get("day_block_size", 7))
    max_entries = int(profile.get("max_entries_per_day", block_size - 1))
    date_col = profile.get("date_column", "D")
    entry_cols = profile["entry_columns"]
    daily = profile.get("daily_total", {})
    formats = profile.get("formats", {})

    month_start = period_start.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
    next_month = (month_start.replace(day=28) + timedelta(days=4)).replace(day=1)
    days_in_month = (next_month - timedelta(days=1)).day

    _set(ws, profile.get("resource_name_cell"), employee_name)
    month_cell = profile.get("month_cell")
    if month_cell:
        ws[month_cell] = month_start
        ws[month_cell].number_format = "mmmm yyyy"

    grouped: dict[Any, list[NormalizedTimeEntry]] = {}
    for entry in entries:
        grouped.setdefault(entry.start.date(), []).append(entry)

    generated_seconds = 0
    for day in range(1, 32):
        header_row = first_day_row + (day - 1) * block_size
        entry_rows = range(header_row + 1, header_row + 1 + max_entries)

        # Clear generated cells but keep all template styling.
        ws[f"{date_col}{header_row}"].value = None
        ws[f"{daily.get('label_column', 'G')}{header_row}"].value = None
        ws[f"{daily.get('value_column', 'H')}{header_row}"].value = None
        for row in entry_rows:
            for col in entry_cols.values():
                ws[f"{col}{row}"].value = None

        if day > days_in_month:
            continue

        current = month_start.replace(day=day)
        current_date = current.date()
        day_entries = grouped.get(current_date, [])
        day_entries = sorted(day_entries, key=lambda item: item.start)
        if len(day_entries) > max_entries:
            raise ReportGenerationError(
                f"{current_date} has {len(day_entries)} entries but template supports only {max_entries} entries per day"
            )

        date_cell = ws[f"{date_col}{header_row}"]
        date_cell.value = current_date
        date_cell.number_format = formats.get("day", "mmmm d, yyyy")

        work_seconds = sum(item.duration_seconds for item in day_entries if not _is_break(item))
        label_cell = ws[f"{daily.get('label_column', 'G')}{header_row}"]
        value_cell = ws[f"{daily.get('value_column', 'H')}{header_row}"]
        is_weekend = current.weekday() >= 5
        if is_weekend and not day_entries:
            value_cell.value = daily.get("weekend_label", "Weekend")
        else:
            label_cell.value = daily.get("work_label", "Work Total")
            value_cell.value = work_seconds / 86400
            value_cell.number_format = formats.get("duration", "[h]:mm")

        for offset, entry in enumerate(day_entries, 1):
            row = header_row + offset
            ws.row_dimensions[row].hidden = False
            is_break = _is_break(entry)
            ws[f"{entry_cols['start_time']}{row}"] = entry.start.replace(tzinfo=None).time()
            ws[f"{entry_cols['end_time']}{row}"] = entry.end.replace(tzinfo=None).time()
            ws[f"{entry_cols['description']}{row}"] = _task_text(entry)
            ws[f"{entry_cols['type']}{row}"] = "Break" if is_break else "Work"
            ws[f"{entry_cols['duration']}{row}"] = entry.duration_seconds / 86400
            ws[f"{entry_cols['start_time']}{row}"].number_format = formats.get("time", "h:mm AM/PM")
            ws[f"{entry_cols['end_time']}{row}"].number_format = formats.get("time", "h:mm AM/PM")
            ws[f"{entry_cols['duration']}{row}"].number_format = formats.get("duration", "[h]:mm")

        # The template keeps unused detail rows collapsed. Reveal only rows that
        # contain generated entries so the emailed workbook is immediately readable.
        for offset, row in enumerate(entry_rows, 1):
            ws.row_dimensions[row].hidden = offset > len(day_entries)

        generated_seconds += work_seconds

    grand_total = profile.get("grand_total_cell")
    if grand_total:
        ws[grand_total] = generated_seconds / 86400
        ws[grand_total].number_format = formats.get("duration", "[h]:mm")

    _write_raw_sheet(workbook, entries, profile)
    return generated_seconds


def generate_report(
    template_path: str | Path,
    output_path: str | Path,
    entries: list[NormalizedTimeEntry],
    profile: dict[str, Any],
    employee_name: str,
    project_name: str,
    period_start: datetime,
    period_end: datetime,
) -> ValidationResult:
    template_path = Path(template_path)
    output_path = Path(output_path)
    if not template_path.exists():
        raise ReportGenerationError(f"Template not found: {template_path}")

    keep_vba = output_path.suffix.lower() == ".xlsm"
    workbook = load_workbook(template_path, keep_vba=keep_vba)
    renderer = profile.get("renderer", "tabular")
    if renderer == "tabular":
        generated_seconds = _generate_tabular(
            workbook, entries, profile, employee_name, project_name, period_start, period_end
        )
    elif renderer == "clockify_import":
        generated_seconds, _ = _generate_import_linked_template(
            workbook, entries, profile, employee_name, project_name, period_start, period_end
        )
    elif renderer == "clockify_helper_calendar":
        generated_seconds, total_cell = _generate_clockify_helper_calendar(
            workbook, entries, profile, employee_name, period_start
        )
    elif renderer == "calendar_timesheet":
        generated_seconds = _generate_calendar_timesheet(workbook, entries, profile, employee_name, period_start)
        total_cell = None
    else:
        raise ReportGenerationError(f"Unknown report renderer: {renderer}")

    generated_formula_cells = (
        {(profile.get("sheet", "Auto Timesheet"), total_cell)}
        if renderer == "clockify_helper_calendar" and total_cell
        else set()
    )
    _finalize_presentation_only(workbook, profile, generated_formula_cells)

    # Ask spreadsheet apps to calculate ReportFlow-generated formulas on open.
    try:
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.calculation.calcMode = "auto"
    except Exception:
        pass

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)

    # Validation must count the same rows the renderer counts. These renderers
    # exclude break-project rows from their work totals, so the source total must too.
    # (The tabular renderer was missing here, which caused source > generated by
    # exactly the total break time.)
    if renderer in {"tabular", "calendar_timesheet", "clockify_helper_calendar"}:
        source_seconds = sum(item.duration_seconds for item in entries if not _is_break(item))
    else:
        source_seconds = sum(item.duration_seconds for item in entries)
    difference = generated_seconds - source_seconds
    result = ValidationResult(
        valid=difference == 0,
        source_seconds=source_seconds,
        generated_seconds=generated_seconds,
        difference_seconds=difference,
        entry_count=len(entries),
    )
    if not result.valid:
        raise ReportGenerationError(f"Validation failed: source={source_seconds}s generated={generated_seconds}s")
    return result


# Backward-compatible name used by Milestone 1 endpoints/tests.
def generate_xlsx(*args, **kwargs) -> ValidationResult:
    return generate_report(*args, **kwargs)
