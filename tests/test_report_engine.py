from datetime import datetime, timezone
from pathlib import Path
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from app.demo import create_demo_template, demo_entries
from app.models import NormalizedTimeEntry
from app.reporting.engine import generate_report, generate_xlsx
from app.reporting.profile import DEFAULT_PROFILE
from app.generate_test_report import resolve_profile


def test_demo_report_generation(tmp_path):
    template = tmp_path / "template.xlsx"
    output = tmp_path / "output.xlsx"
    create_demo_template(template)
    entries = demo_entries()

    result = generate_xlsx(
        template_path=template,
        output_path=output,
        entries=entries,
        profile=DEFAULT_PROFILE,
        employee_name="Demo User",
        project_name="Demo Project",
        period_start=datetime(2026, 8, 17),
        period_end=datetime(2026, 8, 21, 23, 59, 59),
    )

    assert output.exists()
    assert result.valid is True
    assert result.entry_count == 5
    assert result.source_seconds == result.generated_seconds

    wb = load_workbook(output, data_only=False)
    ws = wb["Timesheet"]
    assert ws["B2"].value == "Demo User"
    assert ws["B3"].value == "Demo Project"
    assert ws["C8"].value == "Website maintenance and reporting"
    assert int(ws["F5"].value.total_seconds()) == result.source_seconds


def test_cloud9_monthly_current_template_generation(tmp_path):
    template = tmp_path / "cloud9.xlsx"
    output = tmp_path / "cloud9-output.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Timesheet"
    ws.column_dimensions.group("C", "F")
    ws.column_dimensions["C"].width = 17
    ws.merge_cells("A2:B2")
    ws.merge_cells("A3:B3")
    ws.merge_cells("C2:E2")
    ws.merge_cells("C3:E3")
    ws["A2"] = "Project Manager"
    ws["A3"] = "Department / Project"
    ws["G14"] = "=IF(ISNUMBER('Clockify Import'!A12),'Clockify Import'!D12,IF(AND(NOT(ISNUMBER('Clockify Import'!A12)),ISNUMBER('Clockify Import'!A11)),SUM($G$9:G13),\"\"))"
    ws["G9"] = "=IF(ISNUMBER('Clockify Import'!A7),'Clockify Import'!D7,\"\")"
    ws["G10"] = "=IF(ISNUMBER('Clockify Import'!A8),'Clockify Import'!D8,IF(AND(NOT(ISNUMBER('Clockify Import'!A8)),ISNUMBER('Clockify Import'!A7)),SUM($G$9:G9),\"\"))"
    ws["L2"] = "=IFERROR(DATE(VALUE(MID('Clockify Import'!$B$2,7,4)),VALUE(MID('Clockify Import'!$B$2,4,2)),VALUE(LEFT('Clockify Import'!$B$2,2))),\"\")"
    ws["L3"] = "=IFERROR(DATE(VALUE(RIGHT('Clockify Import'!$B$2,4)),VALUE(MID('Clockify Import'!$B$2,17,2)),VALUE(MID('Clockify Import'!$B$2,14,2))),\"\")"
    for column, value in enumerate(
        ("Sr. No.", "Resource Name", "Date", "Time", "Date", "Time", "Duration (H:MM)", "Task", "Description / Comments"),
        start=1,
    ):
        ws.cell(7, column).value = value
    import_ws = wb.create_sheet("Clockify Import")
    import_ws.sheet_state = "hidden"
    import_ws["B2"] = "01/08/2026 - 31/08/2026"
    import_ws["A7"] = datetime(2026, 8, 1)
    import_ws["B7"] = "Stale description"
    import_ws["C7"] = "Stale project"
    wb.save(template)

    entries = demo_entries()
    profile = resolve_profile(template, "monthly")
    result = generate_xlsx(
        template_path=template,
        output_path=output,
        entries=entries,
        profile=profile,
        employee_name="Demo User",
        project_name="Demo Project",
        period_start=datetime(2026, 8, 1),
        period_end=datetime(2026, 9, 1),
    )

    generated = load_workbook(output, data_only=False)
    generated_ws = generated["Timesheet"]
    first = entries[0]
    assert result.valid is True
    assert generated.sheetnames == ["Timesheet"]
    assert generated_ws["A2"].value == "Project Manager"
    assert generated_ws["A3"].value == "Department / Project"
    assert generated_ws["C2"].value == "Demo User"
    assert generated_ws["C3"].value == "Demo Project"
    assert generated_ws["A9"].value == 1
    assert generated_ws["B9"].value == first.user_name
    assert generated_ws["C9"].value.date() == first.start.date()
    assert generated_ws["D9"].value == first.start.time()
    assert generated_ws["E9"].value.date() == first.end.date()
    assert generated_ws["F9"].value == first.end.time()
    assert generated_ws["G9"].value.total_seconds() == first.duration_seconds
    assert generated_ws["H9"].value == first.description
    assert generated_ws["I9"].value == first.project_name
    assert generated_ws["L2"].value.date() == datetime(2026, 8, 1).date()
    assert generated_ws["L3"].value.date() == datetime(2026, 8, 31).date()
    assert generated_ws["F5"].value is None
    column_ranges = [
        (dimension.min, dimension.max)
        for dimension in generated_ws.column_dimensions.values()
        if dimension.min is not None and dimension.max is not None
    ]
    assert not any(
        max(left[0], right[0]) <= min(left[1], right[1])
        for index, left in enumerate(column_ranges)
        for right in column_ranges[index + 1:]
    )
    assert not any(cell.data_type == "f" for row in generated_ws.iter_rows() for cell in row)
    assert {str(merged) for merged in generated_ws.merged_cells.ranges} >= {
        "A2:B2", "A3:B3", "C2:E2", "C3:E3"
    }


def test_weekly_profile_resolution_remains_detailed_report(tmp_path):
    template = tmp_path / "weekly.xlsx"
    wb = Workbook()
    wb.active.title = "Detailed Report"
    wb.save(template)

    profile = resolve_profile(template, "weekly")

    assert profile["sheet"] == "Detailed Report"
    assert profile["data_start_row"] == 4


def test_atm_link_helper_monthly_writes_values_in_template_calendar_format(tmp_path):
    template = tmp_path / "atm-helper.xlsm"
    output = tmp_path / "atm-helper-output.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Auto Timesheet"
    ws["D5"] = datetime(2026, 6, 1)
    ws["E3"] = '=TEXT(D5, "mmmm yyyy")'
    ws["G5"] = "=SUM('Clockify Helper'!F2:F1000)"
    ws["D6"] = "=INDEX('Clockify Helper'!B2:B1000,1)"
    ws.row_dimensions[26].height = 25.5
    ws["D26"].font = Font(name="Arial", size=10, bold=False)
    ws["G26"].font = Font(name="Arial", size=10, bold=False)
    ws["H26"].font = Font(name="Arial", size=11, bold=False)
    ws.merge_cells("D229:G229")
    ws["D229"] = "Grand Total"
    raw = wb.create_sheet("Clockify Raw")
    for column, header in enumerate(
        ("Project", "Client", "Description", "Task", "User", "Group", "Email", "Tags", "Billable",
         "Start Date", "Start Time", "End Date", "End Time", "Duration (h)", "Duration (decimal)"),
        1,
    ):
        raw.cell(1, column).value = header
    helper = wb.create_sheet("Clockify Helper")
    helper.sheet_state = "hidden"
    helper["A2"] = "=IF('Clockify Raw'!J2=\"\",\"\",'Clockify Raw'!J2)"
    helper["F2"] = "=IF('Clockify Raw'!O2=\"\",\"\",'Clockify Raw'!O2/24)"
    wb.save(template)

    entries = [NormalizedTimeEntry(
        id="weekend-work",
        user_name="Demo User",
        project_name="ATM Link Newsletter Management",
        task_name="Weekend edit",
        description="Saturday work",
        start=datetime(2026, 9, 26, 9, 0, tzinfo=timezone.utc),
        end=datetime(2026, 9, 26, 10, 20, tzinfo=timezone.utc),
        duration_seconds=4800,
    )]
    entries.append(NormalizedTimeEntry(
        id="break-entry",
        user_name="Demo User",
        project_name="Break",
        description="Break",
        start=datetime(2026, 9, 26, 10, 20, tzinfo=timezone.utc),
        end=datetime(2026, 9, 26, 10, 40, tzinfo=timezone.utc),
        duration_seconds=1200,
    ))
    for offset in range(4):
        start = datetime(2026, 9, 26, 11 + offset, 0, tzinfo=timezone.utc)
        entries.append(NormalizedTimeEntry(
            id=f"extra-weekend-{offset}",
            user_name="Demo User",
            project_name="ATM Link Newsletter Management",
            task_name=f"Weekend task {offset + 1}",
            description=f"Weekend work {offset + 1}",
            start=start,
            end=start.replace(minute=20),
            duration_seconds=1200,
        ))
    profile = resolve_profile(template, "monthly")
    result = generate_xlsx(
        template_path=template,
        output_path=output,
        entries=entries,
        profile=profile,
        employee_name="Demo User",
        project_name="ATM Link",
        period_start=datetime(2026, 9, 1),
        period_end=datetime(2026, 10, 1),
    )

    generated = load_workbook(output, data_only=False)
    generated_ws = generated["Auto Timesheet"]
    first = entries[0]
    assert profile["renderer"] == "clockify_helper_calendar"
    assert result.valid is True
    assert result.generated_seconds == sum(
        entry.duration_seconds for entry in entries if entry.project_name.strip().lower() != "break"
    )
    assert generated_ws["D5"].value == "September 1, 2026"
    assert generated_ws["D26"].value == "September 4, 2026"
    assert generated_ws["D180"].value == "September 26, 2026"
    assert generated_ws["D5"].number_format == "mmmm d, yyyy"
    assert generated_ws["D26"].number_format == "mmmm d, yyyy"
    assert generated_ws.row_dimensions[26].height == generated_ws.row_dimensions[5].height
    assert generated_ws["D26"].font.bold == generated_ws["D5"].font.bold
    assert generated_ws["D26"].font.sz == generated_ws["D5"].font.sz
    assert generated_ws["G26"].font.bold == generated_ws["G5"].font.bold
    assert generated_ws["H26"].font.bold == generated_ws["H5"].font.bold
    assert generated_ws["E3"].value == "September 2026"
    assert generated.sheetnames == ["Auto Timesheet"]
    assert int(generated_ws["H180"].value.total_seconds()) == result.generated_seconds
    assert "Work Total" in generated_ws["G180"].value
    assert generated_ws["D181"].value == datetime(2026, 9, 26, 9, 0).time()
    assert generated_ws["D182"].value == datetime(2026, 9, 26, 10, 20).time()
    assert generated_ws.row_dimensions[181].hidden is False
    assert generated_ws.row_dimensions[182].hidden is False
    assert generated_ws["D186"].value == datetime(2026, 9, 26, 14, 0).time()
    assert generated_ws.row_dimensions[186].hidden is False
    assert generated_ws["D188"].value == "September 27, 2026"
    assert generated_ws["H230"].value == "=SUM(" + ",".join(
        f"H{row}" for row in [5, 12, 19, 26, 33, 40, 47, 54, 61, 68, 75, 82, 89, 96, 103, 110,
                               117, 124, 131, 138, 145, 152, 159, 166, 173, 180, 188, 195, 202, 209, 216]
    ) + ")"
    assert "D230:G230" in {str(merged) for merged in generated_ws.merged_cells.ranges}
    visible_formulas = [cell.coordinate for row in generated_ws.iter_rows() for cell in row if cell.data_type == "f"]
    assert visible_formulas == ["H230"]
    assert generated.calculation.fullCalcOnLoad is True
    assert generated.calculation.forceFullCalc is True
    data_only = load_workbook(output, data_only=True)
    assert data_only["Auto Timesheet"]["H230"].value is None
    with ZipFile(output) as archive:
        assert archive.testzip() is None
        assert len(archive.namelist()) > 5
    assert generated_ws["D181"].value == first.start.time()
    assert generated_ws["E181"].value == first.end.time()
    assert generated_ws["F181"].value == first.description
    assert generated_ws["G181"].value == "Work"
    assert int(generated_ws["H181"].value.total_seconds()) == first.duration_seconds
    generated.close()
    wb.close()


def test_output_has_no_references_to_removed_sheets(tmp_path):
    """Excel shows a repair prompt if content points at a sheet ReportFlow deleted."""
    import json
    import zipfile
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import PatternFill
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "Timesheet"
    for column, value in enumerate(
        ("Sr. No.", "Resource Name", "Date", "Time", "Date", "Time", "Duration (H:MM)", "Task", "Description / Comments"),
        start=1,
    ):
        ws.cell(7, column).value = value
    helper = wb.create_sheet("Clockify Import")
    helper.sheet_state = "hidden"
    wb.defined_names["ImportRange"] = DefinedName("ImportRange", attr_text="'Clockify Import'!$A$7:$D$500")
    ws.conditional_formatting.add(
        "A9:I200", FormulaRule(formula=["ISNUMBER('Clockify Import'!$A7)"], fill=PatternFill(bgColor="DDDDDD"))
    )
    # A rule that only points at its own sheet must survive.
    ws.conditional_formatting.add("A9:I200", FormulaRule(formula=["$A9>5"], fill=PatternFill(bgColor="EEEEEE")))
    dv = DataValidation(type="list", formula1="'Clockify Import'!$B$7:$B$50")
    ws.add_data_validation(dv)
    dv.add("H9:H200")
    template = tmp_path / "t.xlsx"
    output = tmp_path / "o.xlsx"
    wb.save(template)

    profile = json.loads((Path(__file__).resolve().parent.parent / "config" / "profiles" / "cloud9_monthly.json").read_text())
    entry = NormalizedTimeEntry(
        project_name="Cloud 9 Flyer", user_name="Mfatir", description="work",
        start=datetime(2026, 9, 2, 17, 30), end=datetime(2026, 9, 2, 20, 0), duration_seconds=9000,
    )
    generate_report(template, output, [entry], profile, "Mfatir", "Cloud9", datetime(2026, 9, 1), datetime(2026, 10, 1))

    archive = zipfile.ZipFile(output)
    for name in archive.namelist():
        assert "Clockify Import" not in archive.read(name).decode("utf-8", "ignore"), name
    assert "$A9&gt;5" in archive.read("xl/worksheets/sheet1.xml").decode("utf-8")
