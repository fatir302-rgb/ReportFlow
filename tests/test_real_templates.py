from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from app.models import NormalizedTimeEntry
from app.reporting.engine import generate_report
from app.reporting.profile import load_profile

ROOT = Path(__file__).resolve().parents[1]


def entry(project, description, start, end, user="Mfatir"):
    return NormalizedTimeEntry(
        user_name=user,
        project_name=project,
        description=description,
        start=start,
        end=end,
        duration_seconds=int((end - start).total_seconds()),
    )


def august_entries():
    return [
        entry("ATM Link Newsletter Management (Creation, Upload, Distribution)", "ATM Discount", datetime(2026,8,21,23,39), datetime(2026,8,22,2,30)),
        entry("Cloud 9 Website", "News Homepage", datetime(2026,8,21,17,30), datetime(2026,8,21,23,39)),
        entry("Murkez (Social Media Post)", "Promotion", datetime(2026,8,18,20,39), datetime(2026,8,19,1,6)),
        entry("ATM Link Newsletter Management (Creation, Upload, Distribution)", "4g Wireless", datetime(2026,8,17,21,49), datetime(2026,8,17,23,54)),
        entry("Cloud 9 Flyer", "", datetime(2026,8,17,17,52), datetime(2026,8,17,20,12)),
        entry("Murkez (Social Media Post)", "Uploading on All Social Platforms", datetime(2026,8,13,20,0), datetime(2026,8,13,20,58)),
        entry("Murkez (Content Revision / Editing)", "14th August Video Revision", datetime(2026,8,13,18,0), datetime(2026,8,13,20,0)),
        entry("Cloud 9 Social Media", "GPT Ads 5", datetime(2026,8,12,23,0), datetime(2026,8,13,2,12)),
        entry("Murkez (Video Editing)", "", datetime(2026,8,12,17,45), datetime(2026,8,12,23,0)),
        entry("Cloud 9 Flyer", "FAQ design", datetime(2026,8,11,23,16), datetime(2026,8,12,2,15)),
        entry("Murkez (Video Editing)", "", datetime(2026,8,11,20,0), datetime(2026,8,11,23,16)),
        entry("Murkez Events", "14th August Video Planning", datetime(2026,8,11,0,0), datetime(2026,8,11,0,49)),
        entry("Cloud 9 Flyer", "Masti Alaska redesign for print", datetime(2026,8,10,19,32), datetime(2026,8,10,22,5)),
        entry("Cloud 9 Print Media", "Luggage tag", datetime(2026,8,6,17,43), datetime(2026,8,6,23,2)),
        entry("Cloud 9 Print Media", "Luggage tag", datetime(2026,8,6,1,11), datetime(2026,8,6,2,7)),
        entry("Taxinvestco", "Logo recreate", datetime(2026,8,5,19,7), datetime(2026,8,5,20,17)),
        entry("Cloud 9 Flyer", "Safari Revision", datetime(2026,8,5,17,52), datetime(2026,8,5,18,11)),
        entry("Cloud 9 Print Media", "Luggage tag", datetime(2026,8,4,22,5), datetime(2026,8,5,0,23)),
        entry("ATM Link Article / Content Publishing (Website, Marketplace, Social Platforms)", "Article Publish on Website, Marketplace, LinkedIn, Facebook (27 Years Anniversary)", datetime(2026,8,4,20,1), datetime(2026,8,4,22,5)),
        entry("ATM Link Newsletter Management (Creation, Upload, Distribution)", "Newsletter Upload (27 Year anniversary)", datetime(2026,8,4,19,30), datetime(2026,8,4,20,0)),
        entry("ATM Link (Content Revision / Editing)", "Emblem and content change Across (Website, Marketplace, Facebook, LinkedIn, Email Signature, Newsletters)", datetime(2026,8,3,20,1), datetime(2026,8,3,22,21)),
    ]


def test_atm_link_weekly_real_format(tmp_path):
    template = ROOT / "templates/atm_link/weekly.xlsx"
    profile = load_profile(ROOT / "config/profiles/atm_link_weekly.json")
    output = tmp_path / "weekly.xlsx"
    entries = august_entries()[:5]

    result = generate_report(
        template_path=template,
        output_path=output,
        entries=entries,
        profile=profile,
        employee_name="Mfatir",
        project_name="ATM LINK",
        period_start=datetime(2026,8,17),
        period_end=datetime(2026,8,23,23,59,59),
    )

    assert result.valid
    assert result.source_seconds == 17 * 3600 + 52 * 60
    wb = load_workbook(output, data_only=False)
    ws = wb["Detailed Report"]
    assert ws["A1"].value == "ATM LINK - Week 34 - Mfatir"
    assert ws["A4"].value == "ATM Link Newsletter Management (Creation, Upload, Distribution)"
    assert ws["B4"].value == "ATM Discount"
    assert round(ws["H4"].value, 2) == 2.85


def test_atm_link_monthly_real_format(tmp_path):
    template = ROOT / "templates/atm_link/monthly.xlsm"
    profile = load_profile(ROOT / "config/profiles/atm_link_monthly.json")
    output = tmp_path / "monthly.xlsx"
    entries = august_entries()

    result = generate_report(
        template_path=template,
        output_path=output,
        entries=entries,
        profile=profile,
        employee_name="Mfatir",
        project_name="ATM LINK",
        period_start=datetime(2026,8,1),
        period_end=datetime(2026,8,31,23,59,59),
    )

    assert result.valid
    assert result.source_seconds == 53 * 3600 + 50 * 60
    wb = load_workbook(output, data_only=False)
    ws = wb["Auto Timesheet"]
    assert wb.sheetnames == ["Auto Timesheet"]
    assert ws["E2"].value == "Mfatir"
    assert ws["E3"].value.month == 8
    assert ws["E3"].value.year == 2026
    # Aug 21 = day block row 145. Entries are chronological within the day.
    assert ws["F146"].value == "News Homepage"
    assert ws["F147"].value == "ATM Discount"
    assert int(ws["H145"].value.total_seconds()) == 9 * 3600
    assert int(ws["H222"].value.total_seconds()) == result.generated_seconds
    assert not any(cell.data_type == "f" for row in ws.iter_rows() for cell in row)
